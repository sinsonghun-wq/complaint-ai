from __future__ import annotations

import asyncio
import codecs
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any

from fastapi.security import APIKeyHeader
import pandas as pd
import psycopg
from openpyxl import load_workbook
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from pypdf import PdfReader

from .ai import CATEGORIES, analyze, embedding, fallback, fingerprint, infer_csv_mapping
from .db import connection, fetch_all, fetch_one
from .security import actor_from_auth, current_account, issue_token, password_hash, public_user, uuid4, verify_password
from .worker_manager import ensure_csv_worker
from .settings import CSV_LLM_CONFIDENCE_THRESHOLD, CSV_MAX_UPLOAD_BYTES, CSV_RULE_OTHER_MIN_CONTENT_CHARS, DOCUMENT_MAX_UPLOAD_BYTES, FILE_STORAGE_DIR, IMPORT_PROGRESS_ROWS, LLM_IMPORT_CONCURRENCY, LLM_IMPORT_ENABLED, MAX_BATCH_SIZE, TESSDATA_DIR, TESSERACT_CMD, WEB_ORIGINS

auth_header = APIKeyHeader(name="Authorization", auto_error=False)
app = FastAPI(title="ComplaintAI FastAPI", version="1.0.0",dependencies=[Depends(auth_header)], lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=WEB_ORIGINS if WEB_ORIGINS != ["*"] else ["*"], allow_credentials=WEB_ORIGINS != ["*"], allow_methods=["*"], allow_headers=["*"])

DEPARTMENT_CATEGORIES = {category: [category] for category in CATEGORIES}
STATUSES = ["접수", "진행중", "완료", "취소"]
STATIC_DIR = Path(__file__).resolve().parents[1] / "frontend" / "dist"
FILE_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
classification_slots = asyncio.Semaphore(1)
classification_tasks: set[asyncio.Task] = set()


class LoginBody(BaseModel):
    username: str
    password: str


class SignupBody(LoginBody):
    display_name: str = ""


class ComplaintBody(BaseModel):
    title: str = ""
    content: str = ""


class BatchBody(BaseModel):
    complaints: list[dict[str, Any]] = Field(default_factory=list)


class PasswordBody(BaseModel):
    password: str


class StatusBody(BaseModel):
    status: str


class ResponseBody(BaseModel):
    content: str


class UpdateComplaintBody(BaseModel):
    title: str = ""
    content: str = ""


class TransferBody(BaseModel):
    category: str


class CancelBody(BaseModel):
    reason: str = Field(default="", max_length=2000)


class CsvMappingBody(BaseModel):
    profile_name: str = ""
    title_column: str = ""
    content_columns: list[str] = Field(default_factory=list)
    response_column: str = ""
    category_column: str = ""
    save_mapping: bool = True


def owner_id(actor: dict[str, Any]) -> str | None:
    return None if actor["role"] == "admin" else actor["owner_id"]


def allowed(actor: dict[str, Any]) -> list[str]:
    return DEPARTMENT_CATEGORIES.get(actor.get("department"), []) if actor["role"] == "admin" else []


def require_department(actor: dict[str, Any]) -> list[str]:
    categories = allowed(actor)
    if not categories:
        raise HTTPException(403, "부서가 지정된 관리자만 사용할 수 있습니다.")
    return categories


def require_user(actor: dict[str, Any]) -> None:
    if actor["role"] != "user":
        raise HTTPException(403, "일반 사용자 민원 작성 기능입니다.")


def require_admin(actor: dict[str, Any]) -> None:
    if actor["role"] != "admin":
        raise HTTPException(403, "부서 관리자 전용 기능입니다.")


def check_password(actor: dict[str, Any], password: str) -> None:
    user = current_account(actor)
    if not verify_password(password, user["password_salt"], user["password_hash"]):
        raise HTTPException(403, "비밀번호가 올바르지 않습니다.")


def vector_literal(vector: list[float] | None) -> str | None:
    return f"[{','.join(str(value) for value in vector)}]" if vector else None


def row_to_complaint(row: dict[str, Any], source_row: int) -> dict[str, Any] | None:
    headers = list(row)
    if "해결명(SOLUTION_CRTR_NAME)" in headers and "분쟁유형명(DISPUTE_TYPE_NAME)" in headers:
        item, dispute, solution = str(row.get("품목명(ITEM_NAME)", "")).strip(), str(row.get("분쟁유형명(DISPUTE_TYPE_NAME)", "")).strip(), str(row.get("해결명(SOLUTION_CRTR_NAME)", "")).strip()
        return {"source_row": source_row, **fallback(" · ".join(filter(None, [item, dispute])), f"{dispute}\n{solution}")} if solution else None
    def column(words: list[str]) -> str | None:
        return next((header for header in headers if any(re.sub(r"[\s_·-]+", "", word).lower() in re.sub(r"[\s_·-]+", "", str(header)).lower() for word in words)), None)
    title_key, content_key = column(["민원 제목", "제목", "title", "subject"]), column(["민원 내용", "민원내용", "신청원인", "원문", "내용", "complaint", "content", "질문"])
    content = str(row.get(content_key, "")).strip() if content_key else ""
    return {"source_row": source_row, **fallback(str(row.get(title_key, "")) if title_key else "", content)} if content else None


TITLE_HEADER_WORDS = ["민원 제목", "민원제목", "제목", "title", "subject", "질문명", "문의제목", "사건명"]
CONTENT_HEADER_WORDS = ["민원 내용", "민원내용", "신청원인", "신청내용", "신청사항", "문의내용", "상세내용", "원문", "내용", "complaint", "content", "질문"]


def csv_schema_signature(headers: list[str]) -> str:
    normalized = "\x1f".join(re.sub(r"\s+", "", header).lower() for header in headers)
    return hashlib.sha256(normalized.encode()).hexdigest()


def csv_preview(path: Path, encoding: str, size: int = 5) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.DictReader(source)
        headers = [header.strip() for header in (reader.fieldnames or []) if header and header.strip()]
        samples = [{key: str(value or "").strip()[:500] for key, value in row.items() if key} for _, row in zip(range(size), reader)]
    return headers, samples


def default_csv_mapping(headers: list[str]) -> dict[str, Any]:
    def column(words: list[str]) -> str:
        return next((header for header in headers if any(re.sub(r"[\s_·-]+", "", word).lower() in re.sub(r"[\s_·-]+", "", header).lower() for word in words)), "")
    if "해결명(SOLUTION_CRTR_NAME)" in headers and "분쟁유형명(DISPUTE_TYPE_NAME)" in headers:
        return {"title_column": "품목명(ITEM_NAME)", "content_columns": ["분쟁유형명(DISPUTE_TYPE_NAME)", "해결명(SOLUTION_CRTR_NAME)"], "response_column": "", "category_column": "", "confidence": 0.95, "reason": "공정거래위원회 상담 사례 헤더를 인식했습니다.", "source": "profile"}
    title_column, content_column = column(TITLE_HEADER_WORDS), column(CONTENT_HEADER_WORDS)
    return {"title_column": title_column, "content_columns": [content_column] if content_column else [], "response_column": "", "category_column": "", "confidence": 0.85 if content_column else 0.0, "reason": "기본 민원 헤더 후보를 인식했습니다." if content_column else "자동으로 민원 본문 열을 찾지 못했습니다.", "source": "heuristic"}


def validate_csv_mapping(mapping: dict[str, Any], headers: list[str], samples: list[dict[str, Any]]) -> dict[str, Any]:
    allowed_headers = set(headers)
    title_column = str(mapping.get("title_column") or "").strip()
    content_columns = [str(value).strip() for value in mapping.get("content_columns", []) if str(value).strip() in allowed_headers]
    response_column = str(mapping.get("response_column") or "").strip()
    category_column = str(mapping.get("category_column") or "").strip()
    if title_column not in allowed_headers:
        title_column = ""
    if response_column not in allowed_headers:
        response_column = ""
    if category_column not in allowed_headers:
        category_column = ""
    content_columns = list(dict.fromkeys(content_columns))
    nonempty = [" ".join(str(row.get(column, "")).strip() for column in content_columns).strip() for row in samples]
    nonempty_ratio = sum(bool(value) for value in nonempty) / max(len(samples), 1)
    average_length = sum(len(value) for value in nonempty) / max(len(nonempty), 1)
    title_average = sum(len(str(row.get(title_column, "")).strip()) for row in samples) / max(len(samples), 1) if title_column else 0
    supplied_confidence = max(0.0, min(1.0, float(mapping.get("confidence", 0))))
    validation_score = min(1.0, nonempty_ratio * 0.55 + min(average_length / 160, 1.0) * 0.35 + (0.10 if not title_column or title_average <= 180 else 0.0))
    return {"title_column": title_column, "content_columns": content_columns, "response_column": response_column, "category_column": category_column, "confidence": round(min(supplied_confidence, validation_score) if supplied_confidence else validation_score, 2), "reason": str(mapping.get("reason") or ""), "source": str(mapping.get("source") or "manual"), "valid": bool(content_columns and nonempty_ratio >= 0.6 and average_length >= 8)}


def csv_row_to_record(row: dict[str, Any], source_row: int, mapping: dict[str, Any]) -> dict[str, Any] | None:
    content_parts = [str(row.get(column, "")).strip() for column in mapping["content_columns"]]
    content = "\n".join(part for part in content_parts if part)
    if not content:
        return None
    title = str(row.get(mapping.get("title_column", ""), "")).strip()
    return {"source_row": source_row, "title": title, "content": content, "source_response": str(row.get(mapping.get("response_column", ""), "")).strip() if mapping.get("response_column") else "", "source_category": str(row.get(mapping.get("category_column", ""), "")).strip() if mapping.get("category_column") else ""}


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _is_complaint_header(values: list[Any]) -> bool:
    headers = [re.sub(r"\s+", "", _cell_text(value)).lower() for value in values]
    # Use exact labels here. 안내 시트의 설명문에 “제목·신청원인”이 있어 부분 일치로는 오탐이 난다.
    has_title = any(value in {"민원제목", "제목", "title", "subject"} for value in headers)
    has_content = any(value in {"민원내용", "신청원인", "원문", "내용", "complaint", "content", "질문"} for value in headers)
    return has_title and has_content


def _records_from_rows(rows: list[tuple[int, list[Any]]], sheet_name: str = "") -> list[dict[str, Any]]:
    """Find a complaint table inside one worksheet and return one record per row."""
    header_at = next((index for index, (_, values) in enumerate(rows[:30]) if _is_complaint_header(values)), None)
    if header_at is None:
        return []
    _, header_values = rows[header_at]
    headers = [_cell_text(value) for value in header_values]
    records: list[dict[str, Any]] = []
    for source_row, values in rows[header_at + 1:]:
        row = {headers[index]: _cell_text(value) for index, value in enumerate(values) if index < len(headers) and headers[index]}
        record = row_to_complaint(row, source_row)
        if record:
            record["source_sheet"] = sheet_name
            records.append(record)
    return records


def spreadsheet_records(path: Path, suffix: str) -> list[dict[str, Any]]:
    """Read every data sheet; a worksheet is never treated as one complaint."""
    records: list[dict[str, Any]] = []
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = [(number, list(values)) for number, values in enumerate(worksheet.iter_rows(values_only=True), start=1)]
                records.extend(_records_from_rows(rows, worksheet.title))
        finally:
            workbook.close()
    else:
        # XLS cannot be opened by openpyxl. Keep the same multi-sheet behaviour through pandas/xlrd.
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        for sheet_name, frame in sheets.items():
            rows = [(number, row) for number, row in enumerate(frame.fillna("").values.tolist(), start=1)]
            records.extend(_records_from_rows(rows, str(sheet_name)))
    return records


def _text_quality(text: str, required_markers: tuple[str, ...] = ()) -> float:
    """Heuristic extraction confidence, not a claim of ground-truth OCR accuracy."""
    compact = re.sub(r"\s+", "", text or "")
    if not compact:
        return 0.0
    readable = sum(character.isalnum() or "가" <= character <= "힣" or character in ".,!?·()[]'\"-:/" for character in compact)
    score = min(len(compact) / 160, 1.0) * 0.45 + (readable / len(compact)) * 0.35
    score += 0.20 * (sum(marker in text for marker in required_markers) / max(len(required_markers), 1))
    return round(min(score, 1.0), 3)


def _ocr_pdf_page(path: Path, page_index: int) -> str:
    """Render only an uncertain page. PyMuPDF/Tesseract remain optional runtime dependencies."""
    try:
        import fitz
        from PIL import Image
        import pytesseract
        document = fitz.open(path)
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        return ocr_image(image, pytesseract)
    except Exception:
        return ""


def _installed_tool(name: str) -> str | None:
    """Find a CLI on PATH, in the active virtual environment, or at LibreOffice's Windows default path."""
    found = shutil.which(name)
    if found:
        return found
    extension = ".exe" if sys.platform == "win32" else ""
    candidates = [Path(sys.executable).with_name(f"{name}{extension}")]
    if name == "soffice" and sys.platform == "win32":
        candidates.extend([
            Path("C:/Program Files/LibreOffice/program/soffice.exe"),
            Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        ])
    if name == "tesseract" and sys.platform == "win32":
        candidates.extend([
            Path("C:/Program Files/Tesseract-OCR/tesseract.exe"),
            Path("C:/Program Files (x86)/Tesseract-OCR/tesseract.exe"),
        ])
    if name == "hwp5txt" and sys.platform == "win32":
        candidates.append(Path(__file__).resolve().parents[1] / ".hwp-parser" / "Scripts" / "hwp5txt.exe")
    return next((str(candidate) for candidate in candidates if candidate.is_file()), None)


def ocr_image(image: Any, pytesseract_module: Any | None = None) -> str:
    """Run Korean/English OCR using ComplaintAI's project-local official language models."""
    if not TESSDATA_DIR.is_dir() or not (TESSDATA_DIR / "kor.traineddata").is_file() or not (TESSDATA_DIR / "eng.traineddata").is_file():
        raise RuntimeError("Tesseract 한국어·영어 언어 데이터가 준비되지 않았습니다.")
    if pytesseract_module is None:
        import pytesseract as pytesseract_module
    executable = TESSERACT_CMD or _installed_tool("tesseract")
    if not executable:
        raise RuntimeError("Tesseract 실행 파일을 찾지 못했습니다.")
    pytesseract_module.pytesseract.tesseract_cmd = executable
    # The project path has no spaces. Do not quote it here: pytesseract passes quotes literally to Tesseract on Windows.
    return pytesseract_module.image_to_string(image, lang="kor+eng", config=f"--tessdata-dir {TESSDATA_DIR}")


def _pdf_page_candidates(path: Path) -> list[tuple[str, float, float]]:
    reader = PdfReader(path)
    native = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
    plumber_text = [""] * len(native)
    image_coverage = [0.0] * len(native)
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages):
                plumber_text[index] = page.extract_text() or ""
                page_area = max(float(page.width * page.height), 1.0)
                image_coverage[index] = min(1.0, sum(float(item.get("width", 0) * item.get("height", 0)) for item in page.images) / page_area)
    except Exception:
        pass

    selected: list[tuple[str, float, float]] = []
    for index, primary in enumerate(native):
        alternate = plumber_text[index]
        primary_score = _text_quality(primary, ("신청원인",))
        alternate_score = _text_quality(alternate, ("신청원인",))
        text, score = (alternate, alternate_score) if alternate_score > primary_score else (primary, primary_score)
        # Image coverage is only a guardrail. Text quality remains the routing decision.
        if score < 0.62 or (image_coverage[index] >= 0.70 and score < 0.82):
            ocr_text = _ocr_pdf_page(path, index)
            ocr_score = _text_quality(ocr_text, ("신청원인",))
            if ocr_score > score:
                text, score = ocr_text, ocr_score
        selected.append((text, score, image_coverage[index]))
    return selected


def _records_from_case_pages(pages: list[tuple[str, float, float]], filename: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for page_number, (text, _score, _coverage) in enumerate(pages, start=1):
        heading = re.search(r"사례\s*(\d{3})", text)
        marker = re.search(r"신청원인\s*", text)
        if not heading or not marker:
            continue
        before = text[:marker.start()].splitlines()
        title_lines = [line.strip() for line in before if line.strip() and not re.search(r"사례\s*\d{3}", line)]
        body = text[marker.end():]
        body = re.split(r"피신청인\s*등의\s*주장|가상\s*민원\s*데이터", body, maxsplit=1)[0].strip()
        title = " ".join(title_lines).strip()
        if title and body:
            records.append({"source_row": page_number, "source_case": heading.group(1), **fallback(title, body)})
    if records:
        return records
    combined = "\n".join(text for text, _, _ in pages).strip()
    return [{"source_row": 1, **fallback(Path(filename).stem, combined)}] if combined else []


def _records_from_hwp_text(text: str, filename: str) -> list[dict[str, Any]]:
    """Split the repeated HWP '분야 / 사례 NNN' blocks into individual complaints."""
    headings = list(re.finditer(r"(?m)^\s*\d{2}\s+.+?/\s*사례\s*(\d{3})\s*$", text))
    records: list[dict[str, Any]] = []
    for index, heading in enumerate(headings):
        block = text[heading.end(): headings[index + 1].start() if index + 1 < len(headings) else len(text)]
        marker = re.search(r"신청원인\s*", block)
        if not marker:
            continue
        title = " ".join(line.strip() for line in block[:marker.start()].splitlines() if line.strip())
        body = re.split(r"피신청인\s*등의\s*주장|가상\s*민원\s*데이터", block[marker.end():], maxsplit=1)[0].strip()
        if title and body:
            records.append({"source_row": int(heading.group(1)), "source_case": heading.group(1), **fallback(title, body)})
    if records:
        return records
    return _records_from_case_pages([(text, _text_quality(text, ("신청원인",)), 0.0)], filename)


def pdf_records(path: Path, filename: str) -> list[dict[str, Any]]:
    return _records_from_case_pages(_pdf_page_candidates(path), filename)


def _hwp_to_pdf_records(path: Path, filename: str) -> list[dict[str, Any]]:
    """Optional fallback for low-confidence HWP extraction on computers with LibreOffice."""
    converter = _installed_tool("soffice")
    if not converter:
        return []
    conversion_root = FILE_STORAGE_DIR / "conversion-temp"
    conversion_root.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory(prefix="complaintai-hwp-", dir=conversion_root) as output_dir:
            converted = subprocess.run([converter, "--headless", "--convert-to", "pdf", "--outdir", output_dir, str(path)], capture_output=True, check=False)
            pdf_path = Path(output_dir) / f"{path.stem}.pdf"
            return pdf_records(pdf_path, filename) if not converted.returncode and pdf_path.exists() else []
    except OSError:
        # A missing converter or an unavailable conversion folder must not turn an HWP setup issue into HTTP 500.
        return []


def hwp_records(path: Path, filename: str) -> list[dict[str, Any]]:
    """Prefer pyhwp text, then convert only low-confidence HWP input for the PDF/OCR pipeline."""
    extractor = _installed_tool("hwp5txt")
    text = ""
    if extractor:
        completed = subprocess.run([extractor, str(path)], capture_output=True, check=False)
        if not completed.returncode:
            text = completed.stdout.decode("utf-8", errors="replace").strip()
    confidence = _text_quality(text, ("신청원인",))
    if text and confidence >= 0.62:
        return _records_from_hwp_text(text, filename)
    converted_records = _hwp_to_pdf_records(path, filename)
    if converted_records:
        return converted_records
    if text:
        return _records_from_hwp_text(text, filename)
    raise HTTPException(503, "HWP 텍스트 추출기(hwp5txt/pyhwp) 또는 HWP-to-PDF 변환기가 설치되어 있지 않습니다.")


def insert_complaint(conn: psycopg.Connection, record: dict[str, Any], source_file: str | None, source_row: int | None, stored_owner: str | None, vector: list[float] | None = None) -> bool:
    metadata = json.dumps({key: record.get(key) for key in ["key_points", "urgency", "needs_review", "review_reason", "reason", "keywords"]}, ensure_ascii=False)
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO complaints (title,content,summary,category,source_file,source_row,content_fingerprint,processing_mode,llm_model,embedding_model,prompt_version,analysis_metadata,embedding,owner_user_id)
            SELECT %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::vector,%s
            WHERE NOT EXISTS (SELECT 1 FROM complaints WHERE content_fingerprint=%s AND deleted_at IS NULL AND complaint_status<>'취소' AND owner_user_id IS NOT DISTINCT FROM %s) RETURNING id""",
            (record["title"], record["content"], record["summary"], record["category"], source_file, source_row, fingerprint(record["content"]), record.get("processing_mode", "fallback"), record.get("model"), "Qwen/Qwen3-Embedding-4B" if vector else None, record.get("prompt_version"), metadata, vector_literal(vector), stored_owner, fingerprint(record["content"]), stored_owner))
        return cur.fetchone() is not None


async def save_records(records: list[dict[str, Any]], actor: dict[str, Any], source_file: str | None = None) -> tuple[int, int]:
    saved = skipped = 0
    with connection() as conn:
        for item in records:
            record = await analyze(str(item.get("title", "")), str(item.get("content", ""))) if item.get("use_ai") else fallback(str(item.get("title", "")), str(item.get("content", "")))
            if item.get("category") in CATEGORIES and not item.get("use_ai"):
                record["category"] = item["category"]
            if not record["content"]:
                continue
            was_saved = insert_complaint(conn, record, item.get("source_file") or source_file, item.get("source_row"), owner_id(actor), await embedding(record))
            saved += int(was_saved); skipped += int(not was_saved)
        conn.commit()
    return saved, skipped


@app.on_event("startup")
async def startup() -> None:
    # 이전 설치 DB도 답변 임시 저장 상태를 바로 사용할 수 있도록 호환 컬럼을 보완한다.
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1")
            cur.execute("ALTER TABLE complaint_responses ADD COLUMN IF NOT EXISTS response_state TEXT NOT NULL DEFAULT 'sent'")
            cur.execute("ALTER TABLE complaint_responses ADD COLUMN IF NOT EXISTS sent_at TIMESTAMPTZ")
            cur.execute("UPDATE complaint_responses SET sent_at=created_at WHERE response_state='sent' AND sent_at IS NULL")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS cancelled_at TIMESTAMPTZ")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS cancelled_by_role TEXT")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS cancellation_reason TEXT")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS parent_complaint_id BIGINT REFERENCES complaints(id) ON DELETE SET NULL")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS previous_context JSONB")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS analysis_state TEXT NOT NULL DEFAULT 'completed'")
            cur.execute("ALTER TABLE complaints ADD COLUMN IF NOT EXISTS analysis_revision INTEGER NOT NULL DEFAULT 0")
            cur.execute("UPDATE complaints SET analysis_state='pending' WHERE analysis_state='processing' AND deleted_at IS NULL AND complaint_status='접수'")
        conn.commit()
    for row in fetch_all("SELECT id,analysis_revision FROM complaints WHERE analysis_state='pending' AND deleted_at IS NULL AND complaint_status='접수'"):
        task = asyncio.create_task(classify_submission(row['id'], row['analysis_revision']))
        classification_tasks.add(task)
        task.add_done_callback(classification_tasks.discard)


@app.on_event("shutdown")
async def stop_classification() -> None:
    tasks = list(classification_tasks)
    for task in tasks:
        task.cancel()
    await asyncio.gather(*tasks, return_exceptions=True)


async def classify_submission(complaint_id: int, revision: int) -> None:
    # An old request must never restore a deleted or replaced submission.
    async with classification_slots:
        with connection() as conn:
            row = conn.execute("""UPDATE complaints SET analysis_state='processing'
                WHERE id=%s AND analysis_revision=%s AND analysis_state='pending'
                AND deleted_at IS NULL AND complaint_status='접수' RETURNING title,content""", (complaint_id, revision)).fetchone()
        if not row:
            return
        try:
            record = await analyze(row['title'], row['content'])
        except Exception:
            record = fallback(row['title'], row['content'])
        # Publish the department before the optional, potentially slow embedding.
        metadata = json.dumps({key: record.get(key) for key in ['key_points', 'urgency', 'needs_review', 'review_reason', 'reason', 'keywords']}, ensure_ascii=False)
        with connection() as conn:
            result = conn.execute("""UPDATE complaints SET summary=%s,category=%s,analysis_state='completed',
                processing_mode=%s,llm_model=%s,prompt_version=%s,analysis_metadata=%s::jsonb
                WHERE id=%s AND analysis_revision=%s AND analysis_state='processing'
                AND deleted_at IS NULL AND complaint_status='접수' RETURNING id""",
                (record['summary'], record['category'], record.get('processing_mode', 'fallback'), record.get('model'), record.get('prompt_version'), metadata, complaint_id, revision)).fetchone()
        if not result:
            return
    try:
        vector = await embedding(record)
    except Exception:
        vector = None
    if vector:
        with connection() as conn:
            conn.execute("""UPDATE complaints SET embedding=%s::vector,embedding_model='Qwen/Qwen3-Embedding-4B'
                WHERE id=%s AND analysis_revision=%s AND analysis_state='completed'
                AND deleted_at IS NULL AND complaint_status<>'취소'""", (vector_literal(vector), complaint_id, revision))


@app.post('/api/complaints', status_code=201)
def submit_own_complaint(body: ComplaintBody, background_tasks: BackgroundTasks, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_user(actor)
    content = body.content.strip()
    if not content:
        raise HTTPException(400, '민원 내용을 입력해 주세요.')
    with connection() as conn:
        content_hash = fingerprint(content)
        conn.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))', (str(actor['owner_id']) + content_hash,))
        existing = conn.execute("""SELECT id,title,content,category,complaint_status,analysis_state,analysis_revision FROM complaints
            WHERE owner_user_id=%s AND content_fingerprint=%s AND deleted_at IS NULL AND complaint_status<>'취소' LIMIT 1""",
            (actor['owner_id'], content_hash)).fetchone()
        if existing:
            return {'complaint': existing, 'duplicates': 1, 'message': '동일한 민원이 이미 접수되어 있습니다.'}
        result = conn.execute("""INSERT INTO complaints(title,content,summary,category,owner_user_id,content_fingerprint,analysis_state,analysis_revision)
            VALUES(%s,%s,'',NULL,%s,%s,'pending',1) RETURNING id,title,content,category,complaint_status,analysis_state,analysis_revision""",
            (body.title.strip() or '제목 없음', content, actor['owner_id'], content_hash)).fetchone()
    background_tasks.add_task(classify_submission, result['id'], result['analysis_revision'])
    return {'complaint': result, 'message': '민원이 접수되었습니다.'}


@app.get("/health")
def health(): return {"status": "ok", "service": "complaintai-fastapi"}


@app.get("/health/database")
def health_database():
    return {"status": "ok", "database": "connected", "pgvector": bool(fetch_one("SELECT 1 FROM pg_extension WHERE extname='vector'"))}


@app.post("/api/auth/login")
def login(body: LoginBody):
    user = fetch_one("SELECT id,owner_id,username,display_name,account_role,department,password_salt,password_hash FROM app_users WHERE username=%s", (body.username.strip(),))
    if not user or not verify_password(body.password, user["password_salt"], user["password_hash"]):
        raise HTTPException(401, "계정 또는 비밀번호가 올바르지 않습니다.")
    return {"token": issue_token(user), "user": public_user(user)}


@app.post("/api/auth/signup", status_code=201)
def signup(body: SignupBody):
    username = body.username.strip()
    if not username.replace("_", "").replace("-", "").isalnum() or not 3 <= len(username) <= 30 or len(body.password) < 8:
        raise HTTPException(400, "계정은 3~30자, 비밀번호는 8자 이상이어야 합니다.")
    salt, digest = password_hash(body.password)
    user = {"id": uuid4(), "owner_id": uuid4(), "username": username, "display_name": body.display_name.strip() or username, "account_role": "user", "department": None}
    try:
        with connection() as conn:
            with conn.cursor() as cur:
                cur.execute("INSERT INTO app_users(id,owner_id,username,email,display_name,password_salt,password_hash,account_role) VALUES(%s,%s,%s,%s,%s,%s,%s,'user')", (user["id"], user["owner_id"], username, f"{username}@complaintai.local", user["display_name"], salt, digest))
            conn.commit()
    except psycopg.errors.UniqueViolation:
        raise HTTPException(400, "이미 사용 중인 계정입니다.")
    return {"token": issue_token(user), "user": public_user(user)}


@app.get("/api/auth/context")
def auth_context(actor: Annotated[dict, Depends(actor_from_auth)]): return {"user": public_user(current_account(actor))}


@app.post("/api/auth/verify-password")
def verify(body: PasswordBody, actor: Annotated[dict, Depends(actor_from_auth)]): check_password(actor, body.password); return {"verified": True}


@app.post("/api/analyze")
async def analyze_endpoint(body: ComplaintBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_user(actor)
    record = await analyze(body.title, body.content)
    similar = []
    vector = await embedding(record)
    if vector:
        similar = fetch_all("SELECT id,title,summary,category,1-(embedding <=> %s::vector) AS similarity FROM complaints WHERE deleted_at IS NULL AND complaint_status<>'취소' AND owner_user_id=%s AND embedding IS NOT NULL ORDER BY embedding <=> %s::vector LIMIT 5", (vector_literal(vector), actor["owner_id"], vector_literal(vector)))
    return {"analysis": record, "similar": similar, "ai": {"llm_model": "qwen2.5:7b-instruct", "embedding_model": "Qwen/Qwen3-Embedding-4B", "vector_dimension": 1536}}


@app.post("/api/complaints/batch", status_code=201)
async def batch(body: BatchBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    if not body.complaints: raise HTTPException(400, "저장할 민원이 없습니다.")
    if actor["role"] == "user" and (len(body.complaints) != 1 or body.complaints[0].get("source_file")):
        raise HTTPException(403, "일반 사용자는 새 민원을 한 건씩 작성할 수 있습니다.")
    saved = duplicates = 0
    for start in range(0, len(body.complaints), MAX_BATCH_SIZE):
        batch_saved, batch_duplicates = await save_records(body.complaints[start:start + MAX_BATCH_SIZE], actor)
        saved += batch_saved
        duplicates += batch_duplicates
    return {"saved": saved, "duplicates": duplicates}


@app.get("/api/complaints/counts")
def counts(actor: Annotated[dict, Depends(actor_from_auth)]):
    if actor["role"] == "admin":
        args, clause = (), "TRUE"
    else:
        args, clause = (actor["owner_id"],), "owner_user_id=%s AND NOT (complaint_status='취소' AND cancelled_by_role IS NOT DISTINCT FROM 'user')"
    rows = fetch_all(f"SELECT category,COUNT(*)::int count FROM complaints WHERE deleted_at IS NULL AND category IS NOT NULL AND {clause} GROUP BY category", args)
    deleted = fetch_one(f"SELECT COUNT(*)::int count FROM complaints WHERE deleted_at IS NOT NULL AND {clause}", args)
    return {"categories": rows, "deleted": deleted["count"]}


@app.get("/api/complaints")
def complaints(deleted: bool = False, category: str | None = None, limit: int = 100, offset: int = 0, source: str = "all", actor: dict = Depends(actor_from_auth)):
    limit = max(1, min(limit, 100)); offset = max(offset, 0)
    category = category or None
    if actor["role"] == "admin": clause, scope = "TRUE", []
    else: clause, scope = "owner_user_id=%s", [actor["owner_id"]]
    where = "deleted_at IS NOT NULL" if deleted else "deleted_at IS NULL"
    if source not in ("all", "user", "file"):
        raise HTTPException(400, "민원 출처는 all, user, file 중 하나여야 합니다.")
    # File provenance is independent of ownership, including legacy imports.
    if source == "user":
        where += " AND NULLIF(BTRIM(source_file), '') IS NULL"
    elif source == "file":
        where += " AND NULLIF(BTRIM(source_file), '') IS NOT NULL"
    if actor["role"] == "user":
        where += " AND NOT (complaint_status='취소' AND cancelled_by_role IS NOT DISTINCT FROM 'user')"
    params: list[Any] = [*scope, category, category, limit, offset]
    response_filter = "" if actor["role"] == "admin" else " AND r.response_state='sent'"
    sql = f"""SELECT id,title,content,summary,category,analysis_state,complaint_status,source_file,source_row,processing_mode,llm_model,embedding_model,created_at,deleted_at,cancelled_at,cancelled_by_role,cancellation_reason,parent_complaint_id,previous_context,
        (NULLIF(BTRIM(source_file), '') IS NULL) AS submitted_by_user,
        COALESCE((SELECT content FROM complaint_responses r WHERE r.complaint_id=complaints.id{response_filter} ORDER BY r.created_at DESC LIMIT 1), '') AS latest_response,
        COALESCE((SELECT response_state FROM complaint_responses r WHERE r.complaint_id=complaints.id{response_filter} ORDER BY r.created_at DESC LIMIT 1), '') AS latest_response_state
        FROM complaints WHERE {where} AND {clause} AND (%s::text IS NULL OR category=%s)
        ORDER BY {'deleted_at' if deleted else 'created_at'} DESC LIMIT %s OFFSET %s"""
    rows = fetch_all(sql, params)
    if actor["role"] == "admin":
        manageable = set(allowed(actor))
        for row in rows:
            row["can_manage"] = row["category"] in manageable
            hide_cancelled_content(row)
    total = fetch_one(f"SELECT COUNT(*)::int count FROM complaints WHERE {where} AND {clause} AND (%s::text IS NULL OR category=%s)", [*scope, category, category])
    return {"complaints": rows, "total": total["count"]}


def scoped_where(actor: dict, parameter: int = 2) -> tuple[str, list[Any]]:
    if actor["role"] == "admin": return " AND category=ANY(%s)", [allowed(actor)]
    return " AND owner_user_id=%s", [actor["owner_id"]]


def hide_cancelled_content(row: dict) -> dict:
    if row.get("complaint_status") == "취소":
        row.update(title="취소된 민원", content="", summary="", latest_response="", latest_response_state="", previous_context=None)
    return row


def locked_complaint(cur, complaint_id: int, actor: dict) -> dict:
    clause, args = scoped_where(actor)
    cur.execute(f"SELECT * FROM complaints WHERE id=%s{clause} FOR UPDATE", (complaint_id, *args))
    row = cur.fetchone()
    if not row:
        raise HTTPException(404, "민원을 찾을 수 없거나 해당 민원에 접근할 수 없습니다.")
    if row['deleted_at']:
        raise HTTPException(409, '해당 민원은 삭제(취소)되었습니다.')
    return row


def require_open_complaint(row: dict) -> None:
    if row["complaint_status"] not in {"접수", "진행중"}:
        raise HTTPException(409, "해당 민원은 삭제(취소)되었습니다." if row["complaint_status"] == "취소" else "답변이 완료된 민원은 변경할 수 없습니다.")


@app.get("/api/complaints/{complaint_id}")
def complaint_detail(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    # 목록과 같은 접근 범위; 관리자는 다른 부서 원문을 읽을 수 있다.
    clause, args = ("", []) if actor["role"] == "admin" else (" AND owner_user_id=%s", [actor["owner_id"]])
    response_filter = "" if actor["role"] == "admin" else " AND r.response_state='sent'"
    row = fetch_one(f"""SELECT id,title,content,summary,category,analysis_state,complaint_status,created_at,deleted_at,cancelled_by_role,cancellation_reason,parent_complaint_id,previous_context,
        (NULLIF(BTRIM(source_file), '') IS NULL) submitted_by_user,
        COALESCE((SELECT content FROM complaint_responses r WHERE r.complaint_id=complaints.id{response_filter} ORDER BY r.created_at DESC LIMIT 1),'') latest_response,
        COALESCE((SELECT response_state FROM complaint_responses r WHERE r.complaint_id=complaints.id{response_filter} ORDER BY r.created_at DESC LIMIT 1),'') latest_response_state
        FROM complaints WHERE id=%s{clause}""", (complaint_id, *args))
    if not row:
        raise HTTPException(404, "민원을 찾을 수 없습니다.")
    row["can_manage"] = actor["role"] == "admin" and row["category"] in allowed(actor)
    if actor["role"] == "admin" or row.get("cancelled_by_role") == "user":
        hide_cancelled_content(row)
    return {"complaint": row}


@app.post("/api/department/complaints/{complaint_id}/start")
def start_complaint(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_department(actor)
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            if row["complaint_status"] == "취소":
                raise HTTPException(409, "해당 민원은 삭제(취소)되었습니다.")
            if row["complaint_status"] == "접수":
                cur.execute("UPDATE complaints SET complaint_status='진행중',status_updated_at=NOW() WHERE id=%s", (complaint_id,))
    return complaint_detail(complaint_id, actor)


@app.patch("/api/complaints/{complaint_id}")
async def update_own_complaint(complaint_id: int, body: UpdateComplaintBody, background_tasks: BackgroundTasks, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_user(actor)
    title, content = body.title.strip(), body.content.strip()
    if not content:
        raise HTTPException(400, "민원 내용을 입력해 주세요.")
    original = fetch_one("SELECT complaint_status FROM complaints WHERE id=%s AND owner_user_id=%s AND deleted_at IS NULL", (complaint_id, actor["owner_id"]))
    if not original:
        raise HTTPException(404, "본인 민원을 찾을 수 없습니다.")
    if original["complaint_status"] != "접수":
        raise HTTPException(409, "접수 대기 상태의 민원만 수정할 수 있습니다.")
    with connection() as conn:
        with conn.cursor() as cur:
            original = locked_complaint(cur, complaint_id, actor)
            if original["complaint_status"] != "접수":
                raise HTTPException(409, "접수 대기 상태의 민원만 수정할 수 있습니다.")
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND response_state='draft'", (complaint_id,))
            cur.execute("""UPDATE complaints SET title=%s,content=%s,summary='',category=NULL,content_fingerprint=%s,
                processing_mode='pending',llm_model=NULL,prompt_version=NULL,analysis_metadata='{}'::jsonb,embedding=NULL,embedding_model=NULL,
                analysis_state='pending',analysis_revision=analysis_revision+1,status_updated_at=NOW()
                WHERE id=%s AND owner_user_id=%s AND deleted_at IS NULL RETURNING id,title,content,category,complaint_status,analysis_state,analysis_revision""",
                (title or '제목 없음', content, fingerprint(content), complaint_id, actor['owner_id']))
            result = cur.fetchone()
        conn.commit()
    if not result:
        raise HTTPException(404, "수정할 본인 민원을 찾지 못했습니다.")
    background_tasks.add_task(classify_submission, result['id'], result['analysis_revision'])
    return {"complaint": result}


@app.delete("/api/complaints/{complaint_id}")
def soft_delete(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)], body: CancelBody | None = None):
    reason = (body.reason if body else "").strip()
    if actor["role"] == "admin" and not reason:
        raise HTTPException(400, "민원을 취소하는 이유를 입력해 주세요.")
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            require_open_complaint(row)
            cur.execute("UPDATE complaints SET complaint_status='취소',cancelled_at=NOW(),cancelled_by_role=%s,cancellation_reason=%s,status_updated_at=NOW() WHERE id=%s", (actor["role"], reason or "민원인이 해당 민원을 취소했습니다.", complaint_id))
            if actor['role'] == 'user':
                cur.execute("UPDATE complaints SET deleted_at=NOW(),analysis_state='cancelled',analysis_revision=analysis_revision+1 WHERE id=%s", (complaint_id,))
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND response_state='draft'", (complaint_id,))
        conn.commit()
    return {"deleted": 1, "complaint_status": "취소", "message": "민원이 취소되었습니다." if actor["role"] == "user" else "민원이 삭제되었습니다."}


@app.post("/api/complaints/{complaint_id}/follow-up", status_code=201)
async def follow_up(complaint_id: int, body: ComplaintBody, background_tasks: BackgroundTasks, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_user(actor)
    if not body.content.strip():
        raise HTTPException(400, "새 민원 내용을 입력해 주세요.")
    parent = fetch_one("SELECT id FROM complaints WHERE id=%s AND owner_user_id=%s AND complaint_status='완료' AND deleted_at IS NULL", (complaint_id, actor["owner_id"]))
    if not parent:
        raise HTTPException(409, "답변이 완료된 본인 민원에만 재민원을 접수할 수 있습니다.")
    with connection() as conn:
        with conn.cursor() as cur:
            parent = locked_complaint(cur, complaint_id, actor)
            if parent["complaint_status"] != "완료":
                raise HTTPException(409, "답변이 완료된 본인 민원에만 재민원을 접수할 수 있습니다.")
            cur.execute("SELECT content FROM complaint_responses WHERE complaint_id=%s AND response_state='sent' ORDER BY sent_at DESC LIMIT 1", (complaint_id,))
            answer = cur.fetchone()
            if not answer:
                raise HTTPException(409, "전송 완료된 답변이 없습니다.")
            context = {"complaint_id": complaint_id, "title": parent["title"], "content": parent["content"], "response": answer["content"], "previous_context": parent.get("previous_context")}
            cur.execute("""INSERT INTO complaints(title,content,summary,category,owner_user_id,complaint_status,parent_complaint_id,previous_context,content_fingerprint,processing_mode,analysis_state,analysis_revision)
                VALUES(%s,%s,'',NULL,%s,'접수',%s,%s::jsonb,%s,'pending','pending',1) RETURNING id,analysis_revision""",
                (body.title.strip() or '제목 없음', body.content.strip(), actor['owner_id'], complaint_id, json.dumps(context, ensure_ascii=False), fingerprint(body.content)))
            result = cur.fetchone()
    background_tasks.add_task(classify_submission, result['id'], result['analysis_revision'])
    return {"complaint": result, "message": "이전 민원과 답변을 포함한 재민원이 접수되었습니다."}


@app.delete("/api/complaints/category/{category}")
def delete_category(category: str, body: CancelBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    if category not in CATEGORIES: raise HTTPException(400, "허용되지 않은 카테고리입니다.")
    if category not in require_department(actor):
        raise HTTPException(403, "담당 부서의 민원만 전체 삭제할 수 있습니다.")
    clause, args = scoped_where(actor)
    if not body.reason.strip():
        raise HTTPException(400, "민원을 취소하는 이유를 입력해 주세요.")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(f"UPDATE complaints SET complaint_status='취소',cancelled_at=NOW(),cancelled_by_role='admin',cancellation_reason=%s,status_updated_at=NOW() WHERE category=%s AND deleted_at IS NULL AND complaint_status IN ('접수','진행중'){clause} RETURNING id", (body.reason.strip(), category, *args))
            ids = [row["id"] for row in cur.fetchall()]; moved = len(ids)
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=ANY(%s) AND response_state='draft'", (ids,))
        conn.commit()
    return {"deleted": moved}


@app.post("/api/complaints/{complaint_id}/restore")
def restore(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    clause, args = scoped_where(actor)
    with connection() as conn:
        with conn.cursor() as cur: cur.execute(f"UPDATE complaints SET deleted_at=NULL WHERE id=%s AND deleted_at IS NOT NULL AND cancelled_at IS NULL{clause} RETURNING id", (complaint_id, *args)); row = cur.fetchone()
        conn.commit()
    return {"restored": int(bool(row))}


@app.delete("/api/complaints/{complaint_id}/permanent")
def permanent(complaint_id: int, body: PasswordBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    check_password(actor, body.password); clause, args = scoped_where(actor)
    with connection() as conn:
        with conn.cursor() as cur: cur.execute(f"DELETE FROM complaints WHERE id=%s AND deleted_at IS NOT NULL AND complaint_status NOT IN ('완료','취소'){clause} RETURNING id", (complaint_id, *args)); row = cur.fetchone()
        conn.commit()
    return {"permanently_deleted": int(bool(row))}


@app.delete("/api/complaints/deleted/all")
def permanent_all(body: PasswordBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    # Temporary cross-department test cleanup; do not grant access to users.
    require_admin(actor)
    check_password(actor, body.password)
    with connection() as conn:
        with conn.cursor() as cur: cur.execute("DELETE FROM complaints WHERE deleted_at IS NOT NULL RETURNING id"); count = len(cur.fetchall())
        conn.commit()
    return {"permanently_deleted": count}


@app.delete("/api/complaints/department/all")
def permanent_department_all(actor: Annotated[dict, Depends(actor_from_auth)]):
    """Testing-only archive of all active complaints, including terminal states."""
    # This route is deliberately restricted to department administrators.  It is
    # a temporary test-reset tool and must be removed before production release.
    require_department(actor)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE complaints SET complaint_status='취소',cancelled_at=NOW(),cancelled_by_role='admin',cancellation_reason='테스트 데이터 전체 정리',status_updated_at=NOW(),deleted_at=NOW() WHERE deleted_at IS NULL RETURNING id")
            ids = [row["id"] for row in cur.fetchall()]
            count = len(ids)
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=ANY(%s) AND response_state='draft'", (ids,))
        conn.commit()
    return {"deleted": count, "scope": "all_categories"}


@app.post("/api/cleanup/deleted")
def cleanup(body: PasswordBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    check_password(actor, body.password); clause, args = scoped_where(actor, 1)
    with connection() as conn:
        with conn.cursor() as cur: cur.execute(f"DELETE FROM complaints WHERE id IN (SELECT id FROM complaints WHERE deleted_at IS NOT NULL AND complaint_status NOT IN ('완료','취소'){clause} ORDER BY deleted_at ASC LIMIT 1000) RETURNING id", args); count = len(cur.fetchall())
        conn.commit()
    return {"permanently_deleted": count}


async def save_upload(upload: UploadFile, maximum: int) -> tuple[Path, str]:
    suffix = Path(upload.filename or "upload").suffix.lower()
    target = FILE_STORAGE_DIR / f"{uuid.uuid4()}{suffix}"
    total = 0
    with target.open("wb") as output:
        while chunk := await upload.read(1024 * 1024):
            total += len(chunk)
            if total > maximum:
                output.close(); target.unlink(missing_ok=True); raise HTTPException(413, "파일 크기 제한을 초과했습니다.")
            output.write(chunk)
    return target, suffix


def csv_encoding(path: Path) -> str:
    # Read a bounded sample; do not load a potentially 1GB CSV into memory.
    with path.open("rb") as source:
        sample = source.read(65536)
    for marker, encoding in ((codecs.BOM_UTF8, "utf-8-sig"),
                             (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
                             (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")):
        if sample.startswith(marker):
            return encoding
    try:
        # An incomplete character at the sample boundary is not invalid UTF-8.
        codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp949"


def csv_rows(path: Path, encoding: str):
    with path.open("r", encoding=encoding, newline="") as source:
        yield from csv.DictReader(source)


async def analyze_csv_batch(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    def rule_result(record: dict[str, Any]) -> dict[str, Any]:
        return {**fallback(record["title"], record["content"]), "source_row": record["source_row"], "processing_mode": "rule"}

    def should_use_llm(record: dict[str, Any], result: dict[str, Any]) -> bool:
        if not LLM_IMPORT_ENABLED or result.get("needs_review"):
            return False
        # A sufficiently detailed complaint without any department signal is
        # safely kept in the catch-all category.  Sending every such row to the
        # LLM would make consumer-style CSV imports impractically slow.
        if result.get("category") == "기타" and len(record["content"].strip()) >= CSV_RULE_OTHER_MIN_CONTENT_CHARS:
            return False
        return float(result.get("confidence") or 0) < CSV_LLM_CONFIDENCE_THRESHOLD

    prepared = [(record, rule_result(record)) for record in records]
    llm_candidates = [(record, result) for record, result in prepared if should_use_llm(record, result)]
    if not llm_candidates:
        return [result for _, result in prepared]
    semaphore = asyncio.Semaphore(LLM_IMPORT_CONCURRENCY)

    async def analyze_one(record: dict[str, Any], rule: dict[str, Any]) -> tuple[int, dict[str, Any]]:
        async with semaphore:
            result = await analyze(record["title"], record["content"])
            # Bulk CSV imports do not retry a 60-second local model timeout.
            # The rule result is retained when the LLM is unavailable.
            return record["source_row"], ({**result, "source_row": record["source_row"]} if result.get("processing_mode") == "llm" else rule)

    llm_results = dict(await asyncio.gather(*(analyze_one(record, rule) for record, rule in llm_candidates)))
    return [llm_results.get(record["source_row"], rule) for record, rule in prepared]


def persist_csv_batch(conn: psycopg.Connection, records: list[dict[str, Any]], source_file: str, stored_owner: str | None) -> tuple[int, int]:
    saved = skipped = 0
    for record in records:
        vector = record.pop("_embedding", None)
        if insert_complaint(conn, record, source_file, record["source_row"], stored_owner, vector):
            saved += 1
        else:
            skipped += 1
    return saved, skipped


def update_import_progress(job_id: str, completed: int, *, heartbeat_only: bool = False) -> None:
    with connection() as conn:
        with conn.cursor() as cur:
            if heartbeat_only:
                cur.execute("UPDATE import_jobs SET heartbeat_at=NOW() WHERE id=%s", (job_id,))
            else:
                cur.execute("UPDATE import_jobs SET completed_rows=%s,heartbeat_at=NOW() WHERE id=%s", (completed, job_id))
        conn.commit()


def process_csv_job(job: dict[str, Any]) -> None:
    """Run one queued CSV job from its last durable 500-row checkpoint.

    The worker owns this function.  Progress is visible every small group, while
    complaint inserts remain committed in MAX_BATCH_SIZE-sized transactions.
    """
    job_id = str(job["id"])
    path = Path(job["storage_path"])
    source_file = job["source_file"]
    stored_owner = job["owner_user_id"]
    encoding = job["encoding"]
    mapping = job["column_mapping"]
    checkpoint = int(job.get("checkpoint_rows") or 0)
    try:
        if not path.is_file():
            raise FileNotFoundError("원본 CSV 파일을 찾지 못했습니다.")
        batch: list[dict[str, Any]] = []
        failures: list[tuple[int, dict[str, Any], str]] = []
        completed = checkpoint
        saved = int(job.get("saved_rows") or 0)
        skipped = int(job.get("skipped_rows") or 0)
        failed = int(job.get("failed_rows") or 0)
        pending: list[tuple[int, dict[str, Any]]] = []

        def flush_batch() -> None:
            nonlocal batch, failures, saved, skipped, failed
            if not batch and not failures:
                return
            with connection() as conn:
                batch_saved, batch_skipped = persist_csv_batch(conn, batch, source_file, stored_owner)
                saved += batch_saved
                skipped += batch_skipped
                with conn.cursor() as cur:
                    for row_number, raw, reason in failures:
                        cur.execute("INSERT INTO import_failures(job_id,source_row,raw_data,reason) VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING", (job_id, row_number, json.dumps(raw, ensure_ascii=False), reason))
                    failed += len(failures)
                    cur.execute("UPDATE import_jobs SET checkpoint_rows=%s,completed_rows=%s,saved_rows=%s,skipped_rows=%s,failed_rows=%s,heartbeat_at=NOW() WHERE id=%s", (completed, completed, saved, skipped, failed, job_id))
                conn.commit()
            batch = []
            failures = []

        def process_pending() -> None:
            nonlocal pending, batch, failures, completed
            if not pending:
                return
            for number, row in pending:
                try:
                    record = csv_row_to_record(row, number, mapping)
                    if record:
                        # Keep the heartbeat alive for a slow local LLM request.
                        # Progress is intentionally displayed only every configured
                        # group, but a live worker must not look abandoned mid-group.
                        analyzed = asyncio.run(analyze_csv_batch([record]))[0]
                        analyzed["_embedding"] = asyncio.run(embedding(analyzed))
                        batch.append(analyzed)
                    else:
                        failures.append((number, row, "요약할 민원 내용 열을 찾지 못했습니다."))
                except Exception as error:
                    failures.append((number, row, str(error)))
                completed += 1
                if completed % IMPORT_PROGRESS_ROWS == 0:
                    update_import_progress(job_id, completed)
                else:
                    update_import_progress(job_id, completed, heartbeat_only=True)
            pending = []
            if len(batch) + len(failures) >= MAX_BATCH_SIZE:
                flush_batch()

        for number, row in enumerate(csv_rows(path, encoding), start=2):
            if number <= checkpoint + 1:
                continue
            pending.append((number, row))
            if len(pending) >= IMPORT_PROGRESS_ROWS:
                process_pending()
        process_pending()
        flush_batch()
        with connection() as conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE import_jobs SET status='completed',checkpoint_rows=%s,completed_rows=%s,saved_rows=%s,skipped_rows=%s,failed_rows=%s,completed_at=NOW(),heartbeat_at=NOW(),last_error=NULL WHERE id=%s", (completed, completed, saved, skipped, failed, job_id))
            conn.commit()
    except Exception as error:
        with connection() as conn:
            with conn.cursor() as cur: cur.execute("UPDATE import_jobs SET status='failed',retry_count=retry_count+1,last_error=%s,heartbeat_at=NOW() WHERE id=%s", (str(error), job_id))
            conn.commit()


def start_import_worker(job_id: str) -> None:
    try:
        ensure_csv_worker()
    except Exception as error:
        message = "CSV 처리 워커를 실행하지 못했습니다. 재처리를 눌러 다시 시도하거나 서버 실행 환경을 확인해 주세요."
        with connection() as db:
            with db.cursor() as cur:
                cur.execute("UPDATE import_jobs SET status='failed',last_error=%s WHERE id=%s AND status='queued'", (message, job_id))
            db.commit()
        raise HTTPException(503, message) from error


@app.post("/api/imports", status_code=202)
async def create_import(file: UploadFile = File(...), actor: dict = Depends(actor_from_auth)):
    require_admin(actor)
    path, suffix = await save_upload(file, CSV_MAX_UPLOAD_BYTES)
    if suffix != ".csv": path.unlink(missing_ok=True); raise HTTPException(400, "일괄 처리는 CSV 파일만 지원합니다.")
    try:
        encoding = csv_encoding(path)
        total = sum(1 for _ in csv_rows(path, encoding))
        headers, samples = csv_preview(path, encoding)
    except (UnicodeError, csv.Error) as error:
        path.unlink(missing_ok=True)
        raise HTTPException(422, "CSV 문자 인코딩 또는 형식을 읽을 수 없습니다. 파일을 CSV UTF-8 형식으로 다시 저장해 업로드해 주세요.") from error
    if not headers:
        path.unlink(missing_ok=True); raise HTTPException(422, "CSV 헤더를 찾지 못했습니다.")
    signature = csv_schema_signature(headers)
    stored_mapping = fetch_one("SELECT profile_name,column_mapping,confidence FROM csv_schema_mappings WHERE schema_signature=%s", (signature,))
    if stored_mapping:
        mapping = {**stored_mapping["column_mapping"], "confidence": float(stored_mapping["confidence"]), "source": "saved", "reason": "저장된 CSV 헤더 매핑을 적용했습니다."}
    else:
        llm_mapping = validate_csv_mapping(await infer_csv_mapping(headers, samples), headers, samples)
        heuristic_mapping = validate_csv_mapping(default_csv_mapping(headers), headers, samples)
        mapping = llm_mapping if llm_mapping["valid"] and llm_mapping["confidence"] >= 0.65 else heuristic_mapping
    mapping = validate_csv_mapping(mapping, headers, samples)
    needs_mapping = not mapping["valid"] or mapping["confidence"] < 0.65
    status = "awaiting_mapping" if needs_mapping else "queued"
    job_id = uuid4(); stored_owner = owner_id(actor)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("INSERT INTO source_files(id,original_name,storage_path,mime_type,size_bytes) VALUES(%s,%s,%s,%s,%s)", (job_id, file.filename, str(path), file.content_type, path.stat().st_size))
            cur.execute("INSERT INTO import_jobs(id,source_file,status,total_rows,storage_path,encoding,column_mapping,schema_signature,owner_user_id) VALUES(%s,%s,%s,%s,%s,%s,%s::jsonb,%s,%s)", (job_id, file.filename, status, total, str(path), encoding, json.dumps(mapping, ensure_ascii=False), signature, stored_owner))
            if not needs_mapping and not stored_mapping:
                cur.execute("INSERT INTO csv_schema_mappings(schema_signature,column_mapping,confidence,created_by_user_id) VALUES(%s,%s::jsonb,%s,%s) ON CONFLICT(schema_signature) DO NOTHING", (signature, json.dumps(mapping, ensure_ascii=False), mapping["confidence"], actor["sub"]))
        conn.commit()
    if not needs_mapping:
        await asyncio.to_thread(start_import_worker, str(job_id))
    return {"job_id": job_id, "status": status, "total_rows": total, "batch_size": MAX_BATCH_SIZE, "needs_mapping": needs_mapping, "headers": headers, "samples": samples, "mapping": mapping}


@app.post("/api/imports/{job_id}/mapping", status_code=202)
def confirm_import_mapping(job_id: str, body: CsvMappingBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_admin(actor)
    job = fetch_one("SELECT id,source_file,status,storage_path,encoding,owner_user_id,schema_signature FROM import_jobs WHERE id=%s", (job_id,))
    if not job or job["status"] != "awaiting_mapping":
        raise HTTPException(404, "헤더 매핑 대기 중인 작업을 찾지 못했습니다.")
    path = Path(job["storage_path"])
    if not path.is_file():
        raise HTTPException(404, "원본 CSV 파일을 찾지 못했습니다.")
    headers, samples = csv_preview(path, job["encoding"])
    mapping = validate_csv_mapping({**body.model_dump(), "confidence": 1.0, "source": "manual", "reason": "관리자가 CSV 열 매핑을 확인했습니다."}, headers, samples)
    if not mapping["valid"]:
        raise HTTPException(422, "본문 열을 하나 이상 선택하고, 예시 행에 충분한 민원 내용이 있는지 확인해 주세요.")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE import_jobs SET status='queued',column_mapping=%s::jsonb WHERE id=%s", (json.dumps(mapping, ensure_ascii=False), job_id))
            if body.save_mapping:
                cur.execute("""INSERT INTO csv_schema_mappings(schema_signature,profile_name,column_mapping,confidence,created_by_user_id)
                    VALUES(%s,%s,%s::jsonb,%s,%s)
                    ON CONFLICT(schema_signature) DO UPDATE SET profile_name=EXCLUDED.profile_name,column_mapping=EXCLUDED.column_mapping,confidence=EXCLUDED.confidence,created_by_user_id=EXCLUDED.created_by_user_id,updated_at=NOW()""", (job["schema_signature"], body.profile_name.strip() or None, json.dumps(mapping, ensure_ascii=False), mapping["confidence"], actor["sub"]))
        conn.commit()
    start_import_worker(job_id)
    return {"job_id": job_id, "status": "queued", "mapping": mapping}


@app.get("/api/imports/{job_id}")
def import_status(job_id: str, actor: Annotated[dict, Depends(actor_from_auth)]):
    clause, params = ("", [job_id]) if actor["role"] == "admin" else (" AND owner_user_id=%s", [job_id, actor["owner_id"]])
    row = fetch_one(f"SELECT id,source_file,status,total_rows,completed_rows,checkpoint_rows,saved_rows,skipped_rows,failed_rows,retry_count,last_error,column_mapping,created_at,started_at,heartbeat_at,completed_at FROM import_jobs WHERE id=%s{clause}", params)
    if not row: raise HTTPException(404, "처리 작업을 찾을 수 없습니다.")
    return row


@app.get("/api/imports")
def recent_imports(actor: Annotated[dict, Depends(actor_from_auth)], limit: int = 10):
    require_admin(actor)
    limit = max(1, min(limit, 30))
    jobs = fetch_all("""SELECT id,source_file,status,total_rows,completed_rows,checkpoint_rows,saved_rows,skipped_rows,failed_rows,retry_count,last_error,created_at,started_at,heartbeat_at,completed_at
        FROM import_jobs ORDER BY created_at DESC LIMIT %s""", (limit,))
    return {"jobs": jobs}


@app.post("/api/imports/{job_id}/retry", status_code=202)
def retry_import(job_id: str, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_admin(actor)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("""UPDATE import_jobs
                SET status='queued', completed_rows=checkpoint_rows, heartbeat_at=NULL, worker_id=NULL, last_error=NULL, retry_count=retry_count+1
                WHERE id=%s AND status='failed'
                RETURNING id,status,total_rows,completed_rows,checkpoint_rows,saved_rows,skipped_rows,failed_rows,retry_count,last_error""", (job_id,))
            row = cur.fetchone()
        conn.commit()
    if not row:
        raise HTTPException(404, "재처리할 실패 작업을 찾지 못했습니다.")
    start_import_worker(job_id)
    return row


@app.get("/api/imports/{job_id}/failures")
def import_failures(job_id: str, actor: Annotated[dict, Depends(actor_from_auth)]):
    import_status(job_id, actor)
    return {"failures": fetch_all("SELECT source_row,reason FROM import_failures WHERE job_id=%s ORDER BY source_row LIMIT 100", (job_id,))}


@app.post("/api/intake")
async def intake(file: UploadFile = File(...), actor: dict = Depends(actor_from_auth)):
    require_admin(actor)
    path, suffix = await save_upload(file, DOCUMENT_MAX_UPLOAD_BYTES)
    try:
        records: list[dict[str, Any]] = []
        if suffix == ".csv":
            records = [record for index, row in enumerate(csv_rows(path, csv_encoding(path)), start=2) if (record := row_to_complaint(row, index))]
        elif suffix in {".xlsx", ".xls"}:
            records = spreadsheet_records(path, suffix)
        elif suffix == ".pdf":
            records = pdf_records(path, file.filename or "문서.pdf")
        elif suffix == ".hwp":
            records = hwp_records(path, file.filename or "문서.hwp")
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            try:
                from PIL import Image
                import pytesseract
                text = pytesseract.image_to_string(Image.open(path), lang="kor+eng")
            except Exception as error: raise HTTPException(503, f"OCR을 실행하지 못했습니다: {error}")
            records = [{"source_row": 1, **fallback(Path(file.filename or "이미지").stem, text)}]
        else: raise HTTPException(400, "HWP, PDF, 이미지, XLSX, XLS, CSV 파일만 지원합니다.")
        if not records:
            raise HTTPException(422, "민원 제목과 본문을 가진 데이터를 찾지 못했습니다.")
        for record in records:
            record["source_file"] = file.filename or "업로드 문서"
        return {"file_name": file.filename, "processed": len(records), "max_batch_size": MAX_BATCH_SIZE, "complaints": records}
    except (UnicodeError, csv.Error) as error:
        raise HTTPException(422, "파일 문자 인코딩 또는 형식을 읽을 수 없습니다. CSV는 UTF-8 형식으로 다시 저장해 주세요.") from error
    finally: path.unlink(missing_ok=True)


@app.get("/api/department/context")
def department_context(actor: Annotated[dict, Depends(actor_from_auth)]):
    return {"department": actor.get("department"), "categories": require_department(actor), "statuses": STATUSES}


def department_complaint(complaint_id: int, actor: dict) -> dict:
    row = fetch_one("SELECT id,title,category,complaint_status,deleted_at FROM complaints WHERE id=%s AND category=ANY(%s)", (complaint_id, require_department(actor)))
    if not row: raise HTTPException(404, "소속 부서에서 처리할 수 있는 민원을 찾지 못했습니다.")
    if row["complaint_status"] == "취소" or row['deleted_at']:
        raise HTTPException(409, "해당 민원은 삭제(취소)되었습니다.")
    return row


@app.get("/api/department/complaints")
def department_complaints(status: str = "", actor: dict = Depends(actor_from_auth)):
    categories = require_department(actor)
    rows = fetch_all("""SELECT id,title,content,summary,category,complaint_status,status_updated_at,created_at,(NULLIF(BTRIM(source_file), '') IS NULL) submitted_by_user,
        COALESCE((SELECT content FROM complaint_responses r WHERE r.complaint_id=complaints.id ORDER BY r.created_at DESC LIMIT 1), '') AS latest_response,
        COALESCE((SELECT response_state FROM complaint_responses r WHERE r.complaint_id=complaints.id ORDER BY r.created_at DESC LIMIT 1), '') AS latest_response_state
        FROM complaints WHERE deleted_at IS NULL AND category=ANY(%s) AND (%s='' OR complaint_status=%s) ORDER BY created_at DESC LIMIT 200""", (categories, status, status))
    return {"complaints": [hide_cancelled_content(row) for row in rows]}


@app.patch("/api/department/complaints/{complaint_id}/status")
def department_status(complaint_id: int, body: StatusBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    if body.status != "진행중":
        raise HTTPException(409, "민원 접수 시작, 답변 전송 또는 사유를 입력한 취소 기능으로 상태를 변경해 주세요.")
    return start_complaint(complaint_id, actor)


@app.get("/api/department/complaints/{complaint_id}/responses")
def responses(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    department_complaint(complaint_id, actor)
    return {"responses": fetch_all("SELECT r.id,r.content,r.response_state,r.created_at,r.sent_at,u.display_name author_name FROM complaint_responses r JOIN app_users u ON u.id=r.author_user_id WHERE r.complaint_id=%s ORDER BY r.created_at", (complaint_id,))}


@app.post("/api/department/complaints/{complaint_id}/responses", status_code=201)
def create_response(complaint_id: int, body: ResponseBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    department_complaint(complaint_id, actor)
    content = body.content.strip()
    if not 2 <= len(content) <= 5000: raise HTTPException(400, "응답은 2~5,000자로 작성해 주세요.")
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            require_open_complaint(row)
            cur.execute("SELECT 1 FROM complaint_responses WHERE complaint_id=%s AND response_state='sent' LIMIT 1", (complaint_id,))
            if cur.fetchone():
                raise HTTPException(409, "이미 답변 전송이 완료된 민원입니다.")
            # 같은 관리자가 남긴 이전 임시 답변은 교체해 한 민원에 하나의 최신 초안만 유지한다.
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND author_user_id=%s AND response_state='draft'", (complaint_id, actor["sub"]))
            cur.execute("INSERT INTO complaint_responses(id,complaint_id,author_user_id,department,content,response_state) VALUES(%s,%s,%s,%s,%s,'draft') RETURNING id,content,response_state,created_at", (uuid4(), complaint_id, actor["sub"], actor.get("department"), content)); result = cur.fetchone()
            cur.execute("UPDATE complaints SET complaint_status='진행중',status_updated_at=NOW() WHERE id=%s", (complaint_id,))
        conn.commit()
    return {"response": result, "complaint_status": "진행중"}


@app.delete("/api/department/complaints/{complaint_id}/responses/draft")
def delete_draft_response(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    department_complaint(complaint_id, actor)
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            require_open_complaint(row)
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND author_user_id=%s AND response_state='draft' RETURNING id", (complaint_id, actor["sub"]))
            deleted = len(cur.fetchall())
            if deleted:
                cur.execute("UPDATE complaints SET status_updated_at=NOW() WHERE id=%s", (complaint_id,))
        conn.commit()
    return {"deleted": deleted, "complaint_status": row["complaint_status"]}


@app.post("/api/department/complaints/{complaint_id}/responses/send")
def send_draft_response(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    department_complaint(complaint_id, actor)
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            require_open_complaint(row)
            cur.execute("SELECT 1 FROM complaint_responses WHERE complaint_id=%s AND response_state='sent'", (complaint_id,))
            if cur.fetchone():
                raise HTTPException(409, "이미 답변이 완료된 민원입니다.")
            cur.execute("""UPDATE complaint_responses SET response_state='sent',sent_at=NOW()
                WHERE id=(SELECT id FROM complaint_responses WHERE complaint_id=%s AND author_user_id=%s AND response_state='draft' ORDER BY created_at DESC LIMIT 1)
                RETURNING id,content,response_state,sent_at""", (complaint_id, actor["sub"]))
            response = cur.fetchone()
            if not response:
                raise HTTPException(400, "전송할 임시 저장 답변이 없습니다.")
            cur.execute("UPDATE complaints SET complaint_status='완료',status_updated_at=NOW() WHERE id=%s", (complaint_id,))
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND response_state='draft'", (complaint_id,))
        conn.commit()
    return {"response": response, "complaint_status": "완료"}


@app.post("/api/department/complaints/{complaint_id}/transfer")
def transfer_complaint(complaint_id: int, body: TransferBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_department(actor)
    if body.category not in CATEGORIES:
        raise HTTPException(400, "전달할 부서를 선택해 주세요.")
    department_complaint(complaint_id, actor)
    with connection() as conn:
        with conn.cursor() as cur:
            row = locked_complaint(cur, complaint_id, actor)
            require_open_complaint(row)
            if body.category == row["category"]:
                raise HTTPException(400, "다른 부서를 선택해 주세요.")
            cur.execute("DELETE FROM complaint_responses WHERE complaint_id=%s AND response_state='draft'", (complaint_id,))
            cur.execute("UPDATE complaints SET category=%s,complaint_status='접수',status_updated_at=NOW() WHERE id=%s RETURNING id,category,complaint_status", (body.category, complaint_id))
            result = cur.fetchone()
        conn.commit()
    return {"complaint": result}


@app.get("/api/my/complaints/{complaint_id}/responses")
def my_responses(complaint_id: int, actor: Annotated[dict, Depends(actor_from_auth)]):
    require_user(actor)
    exists = fetch_one("SELECT id FROM complaints WHERE id=%s AND owner_user_id=%s AND complaint_status<>'취소' AND deleted_at IS NULL", (complaint_id, actor["owner_id"]))
    if not exists:
        raise HTTPException(404, "본인 민원을 찾지 못했습니다.")
    return {"responses": fetch_all("SELECT r.id,r.content,r.created_at,r.sent_at,u.display_name author_name FROM complaint_responses r JOIN app_users u ON u.id=r.author_user_id WHERE r.complaint_id=%s AND r.response_state='sent' ORDER BY r.created_at", (complaint_id,))}


@app.get("/{path:path}")
def react_app(path: str, request: Request):
    if path.startswith("api/"):
        raise HTTPException(404, "요청한 API를 찾을 수 없습니다.")
    if STATIC_DIR.exists():
        candidate = STATIC_DIR / path
        if path and candidate.is_file(): return FileResponse(candidate)
        return FileResponse(STATIC_DIR / "index.html")
    raise HTTPException(404, "React 빌드 결과가 없습니다. frontend에서 npm run build를 실행하세요.")

class ChatRequestBody(BaseModel):
    content: str

@app.post("/api/chat", response_model=ChatResponseBody)
async def chatAI(chatBody: ChatRequestBody, actor: Annotated[dict, Depends(actor_from_auth)]):
    user_id = actor["owner_id"]
    # 큐에 요청을 넣고 워커가 처리 완료하여 Future에 결과를 담을 때까지 비동기 대기
    result = await enqueue_chat_request(user_id, chatBody.content)
    return result

